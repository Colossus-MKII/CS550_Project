#!/usr/bin/env python3
"""Build the ACML LaTeX source and editable Word companion from proposal.md.

The official jmlr class compiles the authoritative PDF. The Word document is an
editable single-column companion, not an official ACML Word template.
"""

from __future__ import annotations

import argparse
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE


ROOT = Path(__file__).resolve().parents[1]
LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")


def add_text(paragraph, text: str) -> None:
    """Preserve source links as black, unadorned Word hyperlinks."""
    cursor = 0
    for match in LINK.finditer(text):
        paragraph.add_run(text[cursor : match.start()].replace("**", ""))
        rel_id = paragraph.part.relate_to(
            match.group(2), RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        )
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), rel_id)
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        font = OxmlElement("w:rFonts")
        font.set(qn("w:ascii"), "Times New Roman")
        font.set(qn("w:hAnsi"), "Times New Roman")
        props.append(font)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        props.append(color)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "22")
        props.append(size)
        run.append(props)
        value = OxmlElement("w:t")
        value.text = match.group(1)
        run.append(value)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        cursor = match.end()
    paragraph.add_run(text[cursor:].replace("**", ""))


def page_setup(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    # Companion approximates the official class's 6 x 8.5-inch text block.
    section.top_margin = Inches(1.26)
    section.bottom_margin = Inches(1.24)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.header_distance = Inches(0.8)
    section.footer_distance = Inches(0.25)
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), "1")
    cols.set(qn("w:equalWidth"), "1")


def build(source: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    title = lines[0].removeprefix("# ")
    author = next(line for line in lines if line.startswith("**")).strip(" *")
    document = Document()
    document.settings.odd_and_even_pages_header_footer = False
    document.sections[0].different_first_page_header_footer = False
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
        style = document.styles[name]
        style.font.name = "Times New Roman"
        # Remove template theme overrides so Word and LibreOffice use the
        # explicit conference font, and remove the default decorative title rule.
        fonts = style.element.find(qn("w:rPr")).find(qn("w:rFonts"))
        for attribute in list(fonts.attrib):
            if "theme" in attribute.lower():
                del fonts.attrib[attribute]
        properties = style.element.find(qn("w:pPr"))
        if properties is not None:
            for border in list(properties.findall(qn("w:pBdr"))):
                properties.remove(border)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(13.6)
    document.styles["Title"].font.size = Pt(14.4)
    document.styles["Title"].font.bold = True
    document.styles["Title"].paragraph_format.line_spacing = Pt(18)
    document.styles["Title"].paragraph_format.space_after = Pt(12)
    document.styles["Heading 1"].font.size = Pt(12)
    document.styles["Heading 1"].font.bold = True
    document.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(7)
    document.styles["Heading 1"].paragraph_format.keep_with_next = True
    document.styles["Normal"].paragraph_format.widow_control = True
    settings = document.settings.element
    hyphen = OxmlElement("w:autoHyphenation")
    hyphen.set(qn("w:val"), "true")
    settings.append(hyphen)
    page_setup(document.sections[0])
    header = document.sections[0].header.paragraphs[0]
    header.add_run("CS550 Project Proposal | Fall 2026 | ACML style")
    header.runs[0].font.size = Pt(9)
    paragraph = document.add_paragraph(title, "Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in author.split(", "):
        paragraph = document.add_paragraph(name)
        paragraph.runs[0].bold = True
    paragraph = document.add_paragraph("Rutgers University, CS550 Massive Data Mining")
    paragraph.runs[0].font.size = Pt(10)
    paragraph.runs[0].italic = True
    paragraph.paragraph_format.space_after = Pt(18)
    body_start = lines.index("## Abstract")
    blocks = "\n".join(lines[body_start:]).split("\n\n")
    is_abstract = False
    is_references = False
    for block in blocks:
        text = " ".join(line.strip() for line in block.splitlines()).strip()
        if not text:
            continue
        if text.startswith("## "):
            heading = text.removeprefix("## ")
            is_abstract = heading == "Abstract"
            is_references = heading == "References"
            heading = re.sub(r"^(\d+) ", r"\1. ", heading)
            p = document.add_paragraph(heading, "Heading 1")
            if is_abstract:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
            continue
        if text.startswith("**Keywords:**"):
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Pt(20)
            p.paragraph_format.right_indent = Pt(20)
            p.add_run("Keywords: ").bold = True
            p.add_run(text.removeprefix("**Keywords:** "))
            for run in p.runs:
                run.font.size = Pt(10)
            is_abstract = False
            continue
        if is_references:
            for reference in block.splitlines():
                p = document.add_paragraph()
                p.paragraph_format.left_indent = Pt(11)
                p.paragraph_format.first_line_indent = Pt(-11)
                p.paragraph_format.space_after = Pt(4)
                add_text(p, reference)
        else:
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(0)
            if is_abstract:
                p.paragraph_format.left_indent = Pt(20)
                p.paragraph_format.right_indent = Pt(20)
            add_text(p, text)
            if is_abstract:
                for run in p.runs:
                    run.font.size = Pt(10)
    document.core_properties.title = title
    document.core_properties.author = author
    document.core_properties.subject = "CS550 Fall 2026 final project proposal"
    document.core_properties.keywords = "SAM 3, instance segmentation, tracking, robustness"
    document.core_properties.comments = "Generated from docs/proposal.md"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


CITATIONS = {
    "Aharon et al., 2022": "aharon2022",
    "Luiten et al., 2020": "luiten2020",
    "Meta AI, 2026": "meta2026",
    "Ultralytics, 2026": "ultralytics2026",
    "Yang et al., 2019": "yang2019",
    "Zhang et al., 2022": "zhang2022",
}
REFERENCE_LABELS = [
    ("Aharon et~al.(2022)", "aharon2022"),
    ("Luiten et~al.(2020)", "luiten2020"),
    ("Meta AI(2026)", "meta2026"),
    ("Ultralytics(2026)", "ultralytics2026"),
    ("Yang et~al.(2019)", "yang2019"),
    ("Zhang et~al.(2022)", "zhang2022"),
]


def tex_escape(text: str) -> str:
    escapes = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%",
               "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{",
               "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(escapes.get(character, character) for character in text)


def tex_text(text: str) -> str:
    tokens = re.compile(r"\[[^\]]+\]\(https?://[^)]+\)|\([^()]+\)")
    pieces, cursor = [], 0
    for match in tokens.finditer(text):
        pieces.append(tex_escape(text[cursor:match.start()]))
        link = LINK.fullmatch(match.group())
        if link:
            pieces.append(r"\href{" + link.group(2).replace("%", r"\%") + "}{" + tex_escape(link.group(1)) + "}")
        else:
            keys = match.group()[1:-1].split("; ")
            if all(key in CITATIONS for key in keys):
                pieces.append(r"\citep{" + ",".join(CITATIONS[key] for key in keys) + "}")
            else:
                pieces.append(tex_escape(match.group()))
        cursor = match.end()
    pieces.append(tex_escape(text[cursor:]))
    return "".join(pieces)


def build_latex(source: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    title = tex_escape(lines[0].removeprefix("# "))
    author = next(line for line in lines if line.startswith("**")).strip(" *")
    names = "\\\\\n".join(r"\Name{" + tex_escape(name) + "}" for name in author.split(", "))
    preamble = r"""% Generated from docs/proposal.md; rebuild with tools/build_proposal.py.
% Uses the unmodified class supplied by the official ACML 2026 template.
\documentclass[wcp]{jmlr}
\hypersetup{hypertexnames=false}
\pagenumbering{gobble}
\jmlrproceedings{}{CS550 Project Proposal}
\jmlryear{2026}
\jmlrworkshop{Fall 2026}
\jmlrvolume{}
\jmlrpages{}
\editors{}
% Course metadata only: no claim of publication in ACML/PMLR proceedings.
\makeatletter
\renewcommand*{\@titlefoot}{}
\makeatother
\title[Multi-Object Segmentation and Tracking]{TITLE}
\author[Wu, Guo and Ross]{NAMES\\
\addr Rutgers University, CS550 Massive Data Mining}
\begin{document}
\maketitle
""".replace("TITLE", title).replace("NAMES", names)
    parts = [preamble]
    blocks = "\n".join(lines[lines.index("## Abstract"):]).split("\n\n")
    mode = ""
    reference_index = 0
    for block in blocks:
        text = " ".join(line.strip() for line in block.splitlines()).strip()
        if not text:
            continue
        if text.startswith("## "):
            heading = text.removeprefix("## ")
            if heading == "Abstract":
                mode = "abstract"
                parts.append(r"\begin{abstract}")
            elif heading == "References":
                mode = "references"
                parts.append(r"\begin{thebibliography}{6}")
            else:
                mode = "body"
                parts.append(r"\section{" + tex_escape(re.sub(r"^\d+ ", "", heading)) + "}")
        elif text.startswith("**Keywords:**"):
            parts.append("\n".join([r"\begin{keywords}", tex_text(text.removeprefix("**Keywords:** ")), r"\end{keywords}"]))
        elif mode == "references":
            label, key = REFERENCE_LABELS[reference_index]
            parts.extend([r"\bibitem[" + label + "]{" + key + "}", tex_text(text)])
            reference_index += 1
        else:
            parts.append(tex_text(text))
            if mode == "abstract":
                parts.append(r"\end{abstract}")
    if reference_index != len(REFERENCE_LABELS):
        raise ValueError("proposal reference count does not match citation metadata")
    parts.extend([r"\end{thebibliography}", r"\end{document}"])
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n\n".join(parts) + "\n", encoding="utf-8")
    print(output)


def build_source_zip(latex_source: Path, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    instructions = (
        "CS550 proposal using the official ACML 2026 class.\n\n"
        "Upload this ZIP to Overleaf and set proposal.tex as the main document,\n"
        "or compile with tectonic --untrusted proposal.tex.\n"
        "Standard TeX dependencies come from the compiler/TeX distribution.\n"
        "The class is unmodified. Course metadata replaces proceedings metadata.\n"
        "See TEMPLATE-PROVENANCE.md and LICENSE-LPPL-1.3c.txt.\n"
    )
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(latex_source, "proposal.tex")
        archive.write(ROOT / "templates/acml/jmlr.cls", "jmlr.cls")
        archive.write(ROOT / "templates/acml/LICENSE-LPPL-1.3c.txt", "LICENSE-LPPL-1.3c.txt")
        archive.write(ROOT / "templates/acml/README.md", "TEMPLATE-PROVENANCE.md")
        archive.writestr("README.txt", instructions)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "docs/proposal.md")
    parser.add_argument("--output", type=Path, default=ROOT / "output/docx/proposal.docx")
    parser.add_argument("--latex-output", type=Path, default=ROOT / "paper/proposal.tex")
    parser.add_argument("--source-zip", type=Path, default=ROOT / "output/latex/proposal_acml_source.zip")
    args = parser.parse_args()
    build(args.source, args.output)
    build_latex(args.source, args.latex_output)
    build_source_zip(args.latex_output, args.source_zip)
